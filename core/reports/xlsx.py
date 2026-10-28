"""Excel export (.xlsx via openpyxl). Money cells are numeric (no floats are
used for computation; values are written from Decimal as integers)."""
import io

from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="EEEEEE")
BOLD = Font(bold=True)


def _value(v, kind):
    if v is None or v == "":
        return None
    if kind in ("money", "int"):
        return int(v)
    if kind == "datetime":
        return timezone.localtime(v).replace(tzinfo=None)
    if kind == "rate":
        return float(v)  # display-only percentage; never used for money arithmetic
    return v if kind == "date" else str(v)


def report_xlsx(report, generated_by):
    wb = Workbook()
    ws = wb.active
    ws.title = report.key[:31]
    from core.services.branding import company

    ws.append([f"{company().display_name} — {report.title}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([report.filters_text])
    ws.append([f"{_('Generated')} {timezone.localtime().strftime('%d/%m/%Y %H:%M')} · {generated_by}"])
    ws.append([])
    if report.summary:
        for label, v, kind in report.summary:
            ws.append([str(label), _value(v, kind)])
            if kind == "money":
                ws.cell(ws.max_row, 2).number_format = "#,##0"
        ws.append([])
    widths = {}
    for table in report.tables:
        ws.append([str(table.title)])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=11)
        if not table.columns:
            ws.append([])
            continue
        ws.append([str(c.label) for c in table.columns] + [_("Corrected")])
        for i in range(1, len(table.columns) + 2):
            cell = ws.cell(ws.max_row, i)
            cell.font, cell.fill = BOLD, HEAD
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for r in table.rows:
            ws.append([_value(r.get(c.key), c.kind) for c in table.columns] + [_("yes") if r.get("corrected") else ""])
            for i, c in enumerate(table.columns, start=1):
                cell = ws.cell(ws.max_row, i)
                if c.kind == "money" or c.kind == "int":
                    cell.number_format = "#,##0"
                elif c.kind == "date":
                    cell.number_format = "DD/MM/YYYY"
                elif c.kind == "datetime":
                    cell.number_format = "DD/MM/YYYY HH:MM"
                elif c.kind == "rate":
                    cell.number_format = "0.00%"
                widths[i] = max(widths.get(i, 10), min(48, len(str(r.get(c.key) or "")) + 2))
        tot = table.totals()
        if tot is not None and table.rows:
            ws.append([_("Total") if i == 0 and c.key not in tot else (_value(tot[c.key], c.kind) if c.key in tot
                                                                         else None)
                       for i, c in enumerate(table.columns)])
            for i, c in enumerate(table.columns, start=1):
                ws.cell(ws.max_row, i).font = BOLD
                if c.key in tot:
                    ws.cell(ws.max_row, i).number_format = "#,##0"
        ws.append([])
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 22)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
