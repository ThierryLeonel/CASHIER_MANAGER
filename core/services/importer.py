"""CSV / Excel import for the super administrator.

Flow: upload -> dry run (everything is posted inside a database transaction
that is rolled back, so the preview uses the real business rules and
balances) -> confirm -> the same rows are posted for real. If any row has an
error, nothing is imported.
"""
import csv
import io
from datetime import date, datetime

from django.db import transaction
from django.db.models import Q
from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core.models import (
    Account, BankAccount, CashManager, Counterparty, ExpenseCategory, MicroZone, PaymentMethod, Product, User, Zone,
)
from core.money import MoneyError, to_money
from core.simnumbers import SimNumberError, normalize_sim

from . import operations as ops
from . import permissions as perm
from .errors import BusinessError
from .ledger import Ctx, post

MAX_ROWS = 5000
K = Account.Kind
K_AGENT = Counterparty.Kind.AGENT

TEMPLATES = {
    "customers": (["type", "business_sim", "name", "subzone", "microzone", "master_zone", "other_phone", "notes"],
                  [["AGENT", "691234567", "Agent Ndive", "fako-1", "Molyko", "", "", ""],
                   ["POS", "692345678", "POS Molyko", "fako-1", "Molyko", "", "677000000", ""]]),
    "cash_managers": (["name", "phone", "notes"], [["Cash manager Buea", "650123456", ""],
                                                   ["Mama Rose (no SIM)", "", "phone is optional"]]),
    "opening": (["kind", "identifier", "amount", "date", "note"],
                [["CUSTOMER", "691234567", "150000", "2026-11-01", "owed at launch"],
                 ["CASH_DRAWER", "cashier.fako1", "500000", "2026-11-01", ""],
                 ["UVE", "690000001", "2000000", "2026-11-01", ""],
                 ["CASH_MANAGER", "650123456", "-50000", "2026-11-01", "company owes"],
                 ["SHOP_CASH", "shop-molyko", "300000", "2026-11-01", "shop launch"],
                 ["SHOP_UVE", "shop-molyko", "800000", "2026-11-01", ""],
                 ["SHOP_ER", "shop-molyko", "150000", "2026-11-01", ""]]),
    "shops": (["code", "name", "zone", "sim_number", "address", "active"],
              [["shop-molyko", "Shop Molyko", "fako", "690000099", "Molyko, Buea", "yes"]]),
    "shop_staff": (["username", "first_name", "last_name", "shop", "duties", "password"],
                   [["molyko.om", "Ayuk", "Ebai", "shop-molyko", "OM", "Temp-Pass-2026!"],
                    ["molyko.air", "Grace", "Tabi", "shop-molyko", "AIRTIME, SALES, SERVICES, AGENTS", "Temp-Pass-2026!"],
                    ["molyko.mgr", "Paul", "Ngu", "shop-molyko", "ALL", "Temp-Pass-2026!"]]),
    "shop_items": (["category", "name", "price", "active"],
                   [["Phones", "Smartphone A12", "65000", "yes"], ["Orange energy kits", "Energie kit S", "45000", ""]]),
    "shop_stock": (["shop", "item", "serials", "note"],
                   [["shop-molyko", "Smartphone A12", "356000000000011, 356000000000029", "delivery 12"],
                    ["shop-molyko", "Energie kit S", "OEK-S-0001", ""]]),
    "shop_transactions": (["date", "shop", "type", "amount", "commission", "service", "client_phone", "note"],
                          [["2026-11-02", "shop-molyko", "CASHIN", "50000", "250", "", "677123456", ""],
                           ["2026-11-02", "shop-molyko", "CASHOUT", "20000", "200", "", "", ""],
                           ["2026-11-02", "shop-molyko", "AIRTIME", "5000", "", "", "", ""],
                           ["2026-11-02", "shop-molyko", "GOOD_DEAL", "10000", "", "", "", ""],
                           ["2026-11-02", "shop-molyko", "CONVERT", "100000", "", "", "", ""],
                           ["2026-11-02", "shop-molyko", "SERVICE", "4", "", "OTT", "", "count of services"]]),
    "transactions": (["date", "cashier", "type", "customer_sim", "agent_sim", "product_or_method", "account",
                      "amount", "category", "bank_reference", "note"],
                     [["2026-11-02", "cashier.fako1", "SUPPLY_AGENT", "691234567", "", "UVE", "690000001", "100000",
                       "", "", ""],
                      ["2026-11-02", "cashier.fako1", "SUPPLY_POS", "692345678", "", "CASH", "", "50000", "", "", ""],
                      ["2026-11-02", "cashier.fako1", "PAYMENT", "691234567", "", "CASH", "", "60000", "", "", ""],
                      ["2026-11-02", "cashier.fako1", "PAYMENT", "691234567", "", "BANK", "SAMPLE-0001", "40000", "",
                       "DEP-123", ""],
                      ["2026-11-02", "cashier.fako1", "EXPENSE", "", "", "", "", "5000", "fuel", "", "moto"]]),
}


def template_xlsx(kind):
    head, rows = TEMPLATES[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = kind
    ws.append(head)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(head, start=1):
        ws.cell(1, i).font = Font(bold=True)
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(h) + 4)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_rows(name, content):
    """Return list of dicts with lower-case, stripped header keys."""
    name = (name or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        text = content.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    else:
        raise BusinessError(_("Upload an Excel (.xlsx) or CSV file."))
    rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
    if not rows:
        raise BusinessError(_("The file is empty."))
    head = [str(h or "").strip().lower().replace(" ", "_") for h in rows[0]]
    data = []
    for r in rows[1:]:
        data.append({head[i]: ("" if v is None else v) for i, v in enumerate(r) if i < len(head) and head[i]})
    if len(data) > MAX_ROWS:
        raise BusinessError(_("Too many rows (maximum %(n)s per file).") % {"n": MAX_ROWS})
    return data


def _s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _money(v):
    if isinstance(v, float):
        if not v.is_integer():
            raise BusinessError(_("Amounts must be whole FCFA"))
        v = int(v)
    try:
        return to_money(v)
    except MoneyError as e:
        raise BusinessError(_(str(e)))


def _sim(v, label):
    try:
        return normalize_sim(_s(v))
    except SimNumberError:
        raise BusinessError(_("%(label)s: invalid number %(v)s") % {"label": label, "v": _s(v)})


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = _s(v)
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(t, f).date()
        except ValueError:
            pass
    raise BusinessError(_("Invalid date %(v)s (use YYYY-MM-DD or DD/MM/YYYY)") % {"v": t})


def _microzone(sub, name):
    """Microzone by subzone (code or name) and microzone name; created if it does not exist yet."""
    name, sub = _s(name), _s(sub)
    if not name:
        return None
    if not sub:
        mz = MicroZone.objects.filter(name__iexact=name)
        if mz.count() == 1:
            return mz.first()
        raise BusinessError(_("Give the subzone of microzone %(v)s") % {"v": name})
    z = Zone.objects.filter(parent__isnull=False).filter(Q(code__iexact=sub) | Q(name__iexact=sub)).first()
    if not z:
        raise BusinessError(_("Unknown subzone %(v)s") % {"v": sub})
    return MicroZone.objects.filter(subzone=z, name__iexact=name).first() or MicroZone.objects.create(subzone=z, name=name)


def _zone(v):
    t = _s(v)
    if not t:
        return None
    z = Zone.objects.filter(parent__isnull=True).filter(code__iexact=t).first() or \
        Zone.objects.filter(parent__isnull=True, name__iexact=t).first()
    if not z:
        raise BusinessError(_("Unknown master zone %(v)s") % {"v": t})
    return z


KIND_ALIASES = {"AGENT": "AGENT", "AGENT COMMERCIAL": "AGENT", "COMMERCIAL AGENT": "AGENT", "AG": "AGENT",
                "POS": "POS", "PDV": "POS", "POINT DE VENTE": "POS", "POINT OF SALE": "POS"}


def _customer(sim, kind=None, label=None):
    cp = Counterparty.objects.filter(business_sim=sim).first()
    if not cp:
        raise BusinessError(_("No agent or POS is registered with business SIM %(sim)s. Register it first.")
                            % {"sim": sim})
    if kind and cp.kind != kind:
        raise BusinessError(_("%(cp)s is not a %(kind)s.") % {"cp": cp, "kind": Counterparty.Kind(kind).label})
    return cp


def _cashier(v):
    t = _s(v)
    u = User.objects.filter(role=User.Role.CASHIER).filter(username__iexact=t).first() or \
        User.objects.filter(role=User.Role.CASHIER, staff_code__iexact=t).first()
    if not u:
        raise BusinessError(_("Unknown cashier %(v)s (use the username or staff ID)") % {"v": t})
    return u


def _asset_account(kind, ident, cashier=None):
    ident = _s(ident)
    qs = Account.objects.filter(kind=kind, active=True)
    if kind == K.BANK:
        a = qs.filter(bank__account_number__iexact=ident).first() if ident else qs.first() if qs.count() == 1 else None
    elif kind == K.CASH_DRAWER:
        if ident:
            u = User.objects.filter(username__iexact=ident).first()
            a = (qs.filter(perm.zone_accounts_q(u)).first() if u and u.is_cashier else None) or \
                qs.filter(name__iexact=ident).first() or qs.filter(zone__name__iexact=ident).first()
        else:
            a = qs.filter(perm.zone_accounts_q(cashier)).first() if cashier else None
    else:
        if ident:
            a = qs.filter(sim__msisdn=_sim(ident, kind)).first()
        else:
            mine = list(qs.filter(perm.zone_accounts_q(cashier)).distinct()) if cashier else []
            a = mine[0] if len(mine) == 1 else None
    if not a:
        raise BusinessError(_("Company %(kind)s account not found: %(v)s") % {"kind": Account.Kind(kind).label,
                                                                             "v": ident or "—"})
    return a


# ---------------------------------------------------------------- row handlers
def _row_customer(user, r):
    kind = KIND_ALIASES.get(_s(r.get("type")).upper())
    if not kind:
        raise BusinessError(_("type must be AGENT or POS"))
    sim = _sim(r.get("business_sim"), _("business SIM"))
    name = _s(r.get("name"))
    if not name:
        raise BusinessError(_("name is required"))
    cp = Counterparty.objects.filter(business_sim=sim).first()
    if cp and cp.kind != kind:
        raise BusinessError(_("Business SIM %(sim)s is already registered to %(cp)s.") % {"sim": sim, "cp": cp})
    zone = _zone(r.get("master_zone"))
    mz = _microzone(r.get("subzone"), r.get("microzone"))
    if mz and kind == K_AGENT and mz.agent_id and mz.agent.business_sim != sim:
        raise BusinessError(_("%(mz)s is already controlled by %(agent)s.") % {"mz": mz.name, "agent": mz.agent.name})
    if cp:
        cp.name, cp.master_zone = name, zone or cp.master_zone
        cp.microzone = mz or cp.microzone
        cp.other_phone = _s(r.get("other_phone")) or cp.other_phone
        cp.notes = _s(r.get("notes")) or cp.notes
        cp.save()
        _take_charge(cp)
        return _("updated %(cp)s") % {"cp": cp}
    cp = Counterparty.objects.create(kind=kind, business_sim=sim, name=name, master_zone=zone, microzone=mz,
                                     other_phone=_s(r.get("other_phone")), notes=_s(r.get("notes")), created_by=user)
    _take_charge(cp)
    return _("created %(cp)s") % {"cp": cp}


def _take_charge(cp):
    if cp.kind == K_AGENT and cp.microzone_id and cp.microzone.agent_id is None:
        MicroZone.objects.filter(pk=cp.microzone_id).update(agent=cp)


def _row_cash_manager(user, r):
    phone = _sim(r.get("phone"), _("phone")) if _s(r.get("phone")) else None  # the SIM number is optional
    name = _s(r.get("name"))
    if not name:
        raise BusinessError(_("name is required"))
    cm = (CashManager.objects.filter(phone=phone).first() if phone
          else CashManager.objects.filter(name__iexact=name).first())
    if cm:
        cm.name, cm.notes = name, _s(r.get("notes")) or cm.notes
        cm.save()
        return _("updated %(cp)s") % {"cp": cm}
    cm = CashManager.objects.create(name=name, phone=phone, notes=_s(r.get("notes")))
    return _("created %(cp)s") % {"cp": cm}


def _row_opening(user, r):
    kind = _s(r.get("kind")).upper()
    ident = r.get("identifier")
    if kind in ("CUSTOMER", "AGENT", "POS"):
        acct = _customer(_sim(ident, _("business SIM"))).account
    elif kind == "CASH_MANAGER":
        t = _s(ident)
        cm = (CashManager.objects.filter(phone=_sim(t, _("phone"))).first() if t[:1].isdigit() or t.startswith("+")
              else CashManager.objects.filter(name__iexact=t).first())
        if not cm:
            raise BusinessError(_("Unknown cash manager %(v)s") % {"v": _s(ident)})
        acct = cm.account
    elif kind in ("CASH_DRAWER", "DRAWER", "CASH"):
        acct = _asset_account(K.CASH_DRAWER, ident)
    elif kind in ("UVE", "ER", "MTN", "BANK"):
        acct = _asset_account(kind, ident)
    elif kind in SHOP_KINDS:
        acct = _shop(ident).account(SHOP_KINDS[kind])
    else:
        raise BusinessError(_("kind must be CUSTOMER, CASH_MANAGER, CASH_DRAWER, UVE, ER, MTN, BANK, SHOP_CASH, "
                              "SHOP_UVE or SHOP_ER"))
    d = _date(r.get("date")) if _s(r.get("date")) else None
    t = ops.record_opening_balance(user, account=acct, amount=_money(r.get("amount")), business_date=d,
                                   description=_s(r.get("note")) or _("Imported"))
    return f"{t.reference} {acct}"


TX_TYPES = ("SUPPLY_AGENT", "SUPPLY_POS", "PAYMENT", "EXPENSE")


def _row_transaction(user, r):
    ttype = _s(r.get("type")).upper().replace(" ", "_")
    if ttype not in TX_TYPES:
        raise BusinessError(_("type must be one of %(v)s") % {"v": ", ".join(TX_TYPES)})
    cashier = _cashier(r.get("cashier"))
    ctx = Ctx(entered_by=user, cashier=cashier, day=None, business_date=_date(r.get("date")))
    amount = _money(r.get("amount"))
    pm = _s(r.get("product_or_method")).upper()
    note = _s(r.get("note"))
    if ttype in ("SUPPLY_AGENT", "SUPPLY_POS"):
        if pm not in Product.values:
            raise BusinessError(_("product_or_method must be UVE, ER or CASH"))
        src = _asset_account(ops.PRODUCT_ACCOUNT_KIND[pm], r.get("account"), cashier)
        if ttype == "SUPPLY_AGENT":
            cust, recip = _customer(_sim(r.get("customer_sim"), _("customer_sim")), "AGENT"), None
        else:
            pos = _customer(_sim(r.get("customer_sim"), _("customer_sim")), "POS")
            agent = _customer(_sim(r.get("agent_sim"), "agent_sim"), "AGENT") if _s(r.get("agent_sim")) else None
            cust, recip = (agent, pos) if agent else (pos, None)
        spec = ops.build_supply(ctx, product=pm, customer=cust, recipient=recip, amount=amount, source_account=src,
                                description=note)
    elif ttype == "PAYMENT":
        if pm not in PaymentMethod.values:
            raise BusinessError(_("product_or_method must be CASH, UVE, MTN or BANK"))
        acct = _asset_account(ops.METHOD_ACCOUNT_KIND[pm], r.get("account"), cashier)
        spec = ops.build_collection(ctx, payer=_customer(_sim(r.get("customer_sim"), _("customer_sim"))), method=pm,
                                    amount=amount, account=acct, bank_reference=_s(r.get("bank_reference")),
                                    description=note)
    else:
        code = _s(r.get("category")).lower()
        cat = ExpenseCategory.objects.filter(code__iexact=code).first() or \
            ExpenseCategory.objects.filter(name_en__iexact=code).first() or \
            ExpenseCategory.objects.filter(name_fr__iexact=code).first()
        if not cat:
            raise BusinessError(_("Unknown expense category %(v)s") % {"v": code})
        spec = ops.build_expense(ctx, category=cat, amount=amount,
                                 drawer=_asset_account(K.CASH_DRAWER, r.get("account"), cashier),
                                 description=note or str(cat))
    t = post([spec], ctx)[0]
    return f"{t.reference} {t.display_type} {t.amount}"


# ---------------------------------------------------------------- shops
SHOP_KINDS = {"SHOP_CASH": K.CASH_DRAWER, "SHOP_UVE": K.UVE, "SHOP_ER": K.ER}
YES = ("", "1", "Y", "YES", "OUI", "TRUE", "VRAI", "X")


def _shop(v):
    from core.models import Shop

    t = _s(v)
    shop = Shop.objects.filter(code__iexact=t).first() or Shop.objects.filter(name__iexact=t).first()
    if not shop:
        raise BusinessError(_("Unknown shop %(v)s (use its code or name)") % {"v": t or "—"})
    return shop


def _row_shop(user, r):
    from core.models import Shop

    code, name = _s(r.get("code")).lower(), _s(r.get("name"))
    if not code or not name:
        raise BusinessError(_("code and name are required"))
    zone = None
    if _s(r.get("zone")):
        z = _s(r.get("zone"))
        zone = Zone.objects.filter(Q(code__iexact=z) | Q(name__iexact=z)).first()
        if not zone:
            raise BusinessError(_("Unknown zone %(v)s") % {"v": z})
    sim = _sim(r.get("sim_number"), _("SIM number")) if _s(r.get("sim_number")) else ""
    active = _s(r.get("active")).upper() in YES
    shop = Shop.objects.filter(code__iexact=code).first()
    created = shop is None
    shop = shop or Shop(code=code)
    shop.name, shop.zone, shop.sim_number, shop.address, shop.active = (
        name, zone or shop.zone, sim or shop.sim_number, _s(r.get("address")) or shop.address, active)
    shop.save()
    return (_("created %(cp)s") if created else _("updated %(cp)s")) % {"cp": shop}


def _row_shop_staff(user, r):
    from core.models import SHOP_DUTY_CODES

    from . import shop_admin as sa

    username = _s(r.get("username"))
    shop = _shop(r.get("shop"))
    raw = _s(r.get("duties")).upper().replace(";", ",")
    duties = SHOP_DUTY_CODES if raw in ("", "ALL", "TOUT", "TOUTES") else [d.strip() for d in raw.split(",") if d.strip()]
    bad = [d for d in duties if d not in SHOP_DUTY_CODES]
    if bad:
        raise BusinessError(_("Unknown duties %(v)s (use %(ok)s or ALL)") % {"v": ", ".join(bad),
                                                                           "ok": ", ".join(SHOP_DUTY_CODES)})
    existing = User.objects.filter(username__iexact=username).first()
    if existing:
        if not existing.is_shop_manager:
            raise BusinessError(_("%(u)s is not a shop staff account.") % {"u": username})
        sa.update_staff(user, existing, first_name=_s(r.get("first_name")) or existing.first_name,
                        last_name=_s(r.get("last_name")) or existing.last_name, shop=shop, duties=duties)
        return _("updated %(cp)s") % {"cp": existing.username}
    pw = _s(r.get("password"))
    if not pw:
        raise BusinessError(_("password is required for a new account (it must be changed at the first sign-in)"))
    u, _pw = sa.create_staff(user, username=username, first_name=_s(r.get("first_name")) or username,
                             last_name=_s(r.get("last_name")), shop=shop, duties=duties, password=pw)
    return _("created %(cp)s") % {"cp": u.username}


def _row_shop_item(user, r):
    from core.models import ItemCategory, ShopItem

    cat_name, name = _s(r.get("category")), _s(r.get("name"))
    if not cat_name or not name:
        raise BusinessError(_("category and name are required"))
    cat = ItemCategory.objects.filter(name__iexact=cat_name).first() or ItemCategory.objects.create(name=cat_name)
    price = _money(r.get("price"))
    if price <= 0:
        raise BusinessError(_("price must be greater than zero"))
    item = ShopItem.objects.filter(name__iexact=name).first()
    created = item is None
    item = item or ShopItem(name=name)
    item.category, item.price, item.active = cat, price, _s(r.get("active")).upper() in YES
    item.save()
    return (_("created %(cp)s") if created else _("updated %(cp)s")) % {"cp": item.name}


def _row_shop_stock(user, r):
    from core.models import ShopItem

    from . import shops as sh

    shop = _shop(r.get("shop"))
    name = _s(r.get("item"))
    item = ShopItem.objects.filter(name__iexact=name).first()
    if not item:
        raise BusinessError(_("Unknown item %(v)s (import the items first)") % {"v": name or "—"})
    units = sh.receive_stock(user, shop=shop, item=item, serials=_s(r.get("serials")), note=_s(r.get("note")))
    return _("%(n)s unit(s) of %(i)s received at %(s)s") % {"n": len(units), "i": item.name, "s": shop.name}


SHOP_TX = ("CASHIN", "CASHOUT", "AIRTIME", "GOOD_DEAL", "CONVERT", "SERVICE")


def _row_shop_transaction(user, r):
    from core.models import ShopService

    from . import shops as sh

    ttype = _s(r.get("type")).upper().replace(" ", "_").replace("-", "_")
    ttype = {"CASH_IN": "CASHIN", "CASH_OUT": "CASHOUT", "GOODDEAL": "GOOD_DEAL"}.get(ttype, ttype)
    if ttype not in SHOP_TX:
        raise BusinessError(_("type must be one of %(v)s") % {"v": ", ".join(SHOP_TX)})
    shop, day = _shop(r.get("shop")), _date(r.get("date"))
    note, phone = _s(r.get("note")), _s(r.get("client_phone"))
    if ttype == "SERVICE":
        svc_name = _s(r.get("service"))
        svc = ShopService.objects.filter(kind=ShopService.Kind.COUNT, active=True).filter(
            Q(name__iexact=svc_name) | Q(name__istartswith=svc_name)).first() if svc_name else None
        if not svc:
            raise BusinessError(_("Unknown counted service %(v)s") % {"v": svc_name or "—"})
        n = int(_money(r.get("amount")))
        c = sh.count_service(user, shop=shop, service=svc, count=n, business_date=day, note=note)
        return f"{svc.name} × {c.count}"
    token = sh.BUSINESS_DATE_OVERRIDE.set(day)
    try:
        amount = _money(r.get("amount"))
        com = _money(r.get("commission")) if _s(r.get("commission")) else None
        if ttype == "CASHIN":
            t = sh.cash_in(user, shop=shop, amount=amount, commission=com, client_phone=phone, description=note)
        elif ttype == "CASHOUT":
            t = sh.cash_out(user, shop=shop, amount=amount, commission=com, client_phone=phone, description=note)
        elif ttype in ("AIRTIME", "GOOD_DEAL"):
            t = sh.airtime(user, shop=shop, amount=amount, kind="GOOD_DEAL" if ttype == "GOOD_DEAL" else "NORMAL",
                           client_phone=phone, description=note)
        else:
            t = sh.convert(user, shop=shop, amount=amount, description=note)
    finally:
        sh.BUSINESS_DATE_OVERRIDE.reset(token)
    return f"{t.reference} {t.get_type_display()} {t.amount}"


HANDLERS = {"customers": _row_customer, "cash_managers": _row_cash_manager, "opening": _row_opening,
            "transactions": _row_transaction, "shops": _row_shop, "shop_staff": _row_shop_staff,
            "shop_items": _row_shop_item, "shop_stock": _row_shop_stock,
            "shop_transactions": _row_shop_transaction}


class _Rollback(Exception):
    pass


def run_import(user, kind, rows, *, commit):
    """Returns (results, ok). results: [(row_number, ok, message)]."""
    perm.require_role(user, perm.CONFIG_ROLES, "import")
    if kind not in HANDLERS:
        raise BusinessError(_("Choose what you are importing."))
    handler = HANDLERS[kind]
    results, errors = [], 0
    try:
        with transaction.atomic():
            for i, r in enumerate(rows, start=2):
                try:
                    with transaction.atomic():
                        msg = handler(user, r)
                    results.append((i, True, str(msg)))
                except BusinessError as e:
                    errors += 1
                    results.append((i, False, e.message))
                except Exception as e:  # noqa: BLE001 - shown to the admin row by row
                    errors += 1
                    results.append((i, False, str(e)[:200]))
            if errors or not commit:
                raise _Rollback()
            ops.log_event(user, "import", kind=kind, rows=len(rows))
    except _Rollback:
        pass
    return results, errors == 0
