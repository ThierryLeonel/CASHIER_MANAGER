"""End-to-end HTTP tests: permissions, idempotent forms, receipts, exports, i18n."""
import io
import uuid

import pypdf
from django.core.cache import cache
from django.test import TestCase
from openpyxl import load_workbook

from core.models import Closing, Discrepancy, Transaction, TxType
from core.services import operations as ops
from core.services.closing import closing_accounts

from .helpers import D, OrgMixin

PW = "Sample-Pass-2026!"


def pdf_text(content):
    return "\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(content)).pages)


class WebBase(OrgMixin, TestCase):
    def setUp(self):
        cache.clear()
        self.make_org()
        self.fund_all()

    def client_for(self, username, lang="en"):
        c = self.client_class()
        self.assertTrue(c.login(username=username, password=PW))
        c.cookies["django_language"] = lang
        return c


class AgentIsolation(WebBase):
    def setUp(self):
        super().setUp()
        self.s1 = ops.record_supply(self.c_f1, product="UVE", customer=self.a1, amount=D(70000),
                                    source_account=self.uve_f1)
        self.s2 = ops.record_supply(self.c_f1, product="UVE", customer=self.a2, amount=D(99000),
                                    source_account=self.uve_f1)
        self.agent = self.client_for("agent.ndive")

    def test_agent_sees_only_own_data(self):
        r = self.agent.get("/")
        self.assertContains(r, "70 000")
        self.assertNotContains(r, "99 000")
        # trying to read another agent's statement via the filter is ignored (forced to own account)
        r = self.agent.get("/reports/statement/?preset=all&customer_sim=691000002")
        self.assertContains(r, "70 000")
        self.assertNotContains(r, "99 000")
        x = self.agent.get("/reports/statement/?preset=all&customer_sim=691000002&format=xlsx")
        ws = load_workbook(io.BytesIO(x.content)).active
        values = [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]
        self.assertIn(70000, values)
        self.assertNotIn(99000, values)
        self.assertEqual(self.agent.get(f"/transactions/{self.s2.pk}/").status_code, 404)
        self.assertEqual(self.agent.get(f"/transactions/{self.s2.pk}/receipt.pdf").status_code, 404)
        self.assertEqual(self.agent.get(f"/transactions/{self.s1.pk}/receipt.pdf").status_code, 200)

    def test_agent_prohibited_actions(self):
        for url in ["/transactions/", "/customers/", "/supply/new/", "/collection/new/", "/closing/",
                    "/reports/customers/", "/reports/supplies/", "/setup/opening-balances/",
                    f"/transactions/{self.s1.pk}/correct/", f"/transactions/{self.s1.pk}/transfer-responsibility/"]:
            r = self.agent.get(url)
            self.assertIn(r.status_code, (403, 404), url)
        r = self.agent.post("/supply/new/", {"idem": uuid.uuid4(), "product": "UVE", "customer": "691000001",
                                             "amount": "1000", "source_account": self.uve_f1.pk})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(Transaction.objects.filter(type=TxType.SUPPLY).count(), 2)
        self.assertNotEqual(self.agent.get("/admin/").status_code, 200)


class RolePermissions(WebBase):
    def test_cashier_cannot_correct_or_configure(self):
        t = ops.record_supply(self.c_f1, product="UVE", customer=self.a1, amount=D(1000), source_account=self.uve_f1)
        c = self.client_for("cashier.fako1")
        self.assertEqual(c.get(f"/transactions/{t.pk}/correct/").status_code, 403)
        self.assertEqual(c.post(f"/transactions/{t.pk}/correct/", {"kind": "VOID", "reason": "x"}).status_code, 403)
        self.assertEqual(c.get("/setup/opening-balances/").status_code, 403)
        self.assertEqual(c.get("/setup/conversion-rate/").status_code, 403)
        self.assertNotEqual(c.get("/admin/core/user/").status_code, 200)
        m = self.client_for("manager")
        self.assertEqual(m.get("/supply/new/").status_code, 403)  # managers do not enter operations
        self.assertEqual(m.get("/setup/opening-balances/").status_code, 403)

    def test_cashier_cannot_post_from_unauthorised_wallet(self):
        c = self.client_for("cashier.fako2")
        r = c.post("/supply/new/", {"idem": uuid.uuid4(), "product": "UVE", "customer": "691000001",
                                    "amount": "1000", "source_account": self.uve_f1.pk})
        self.assertEqual(r.status_code, 200)  # form re-rendered with an error
        self.assertFalse(Transaction.objects.filter(type=TxType.SUPPLY).exists())

    def test_cashier_reports_are_scoped(self):
        ops.record_expense(self.c_f1, category=self.fuel(), amount=D(1111), drawer=self.d_f1, description="a")
        ops.record_expense(self.c_f2, category=self.fuel(), amount=D(2222), drawer=self.d_f2, description="b")
        c = self.client_for("cashier.fako1")
        r = c.get(f"/reports/expenses/?preset=all&cashier={self.c_f2.pk}")
        self.assertContains(r, "1 111")
        self.assertNotContains(r, "2 222")
        m = self.client_for("finance")
        r = m.get("/reports/expenses/?preset=all")
        self.assertContains(r, "1 111")
        self.assertContains(r, "2 222")
        r = m.get(f"/reports/expenses/?preset=all&master_zone={self.zones['fako'].pk}&format=xlsx")
        self.assertEqual(r.status_code, 200)

    def fuel(self):
        from core.models import ExpenseCategory

        return ExpenseCategory.objects.get(code="fuel")

    def test_login_lockout(self):
        c = self.client_class()
        for _ in range(5):
            c.post("/login/", {"username": "manager", "password": "wrong"})
        r = c.post("/login/", {"username": "manager", "password": PW})
        self.assertContains(r, "Too many failed attempts")


class WebFlows(WebBase):
    def test_supply_form_double_submit_posts_once(self):
        c = self.client_for("cashier.fako1")
        self.assertEqual(c.get("/supply/new/").status_code, 200)
        data = {"idem": str(uuid.uuid4()), "product": "UVE", "customer": "+237 691 00 00 01", "amount": "100 000",
                "source_account": self.uve_f1.pk, "description": "web"}
        r1 = c.post("/supply/new/", data)
        r2 = c.post("/supply/new/", data)
        self.assertEqual(r1.status_code, 302)
        self.assertEqual(r2.status_code, 302)
        self.assertEqual(Transaction.objects.filter(type=TxType.SUPPLY).count(), 1)
        self.assertEqual(self.cbal(self.a1), D(100000))
        r = c.get(r2["Location"])
        self.assertContains(r, "already recorded")

    def test_validation_messages(self):
        c = self.client_for("cashier.fako1")
        r = c.post("/collection/new/", {"idem": str(uuid.uuid4()), "payer": "699999999", "method": "BANK",
                                        "amount": "12.5", "account": self.bank.pk})
        self.assertContains(r, "No agent or POS is registered")
        self.assertContains(r, "whole FCFA")
        c_fr = self.client_for("cashier.fako1", "fr")
        r = c_fr.post("/collection/new/", {"idem": str(uuid.uuid4()), "payer": "691000001", "method": "BANK",
                                           "amount": "5000", "account": self.bank.pk})
        self.assertContains(r, "Saisissez la référence du versement bancaire")

    def test_french_interface(self):
        c = self.client_for("cashier.fako1", "fr")
        r = c.get("/")
        self.assertContains(r, "Approvisionnement")
        self.assertContains(r, "Clôture")
        r = c.get("/reports/reconciliation/?preset=all")
        self.assertContains(r, "Rapprochement")

    def test_receipt_pdf_and_corrected_stamp(self):
        t = ops.record_collection(self.c_f1, payer=self.a1, method="MTN", amount=D(40000), account=self.mtn_f1)
        c = self.client_for("cashier.fako1")
        r = c.get(f"/transactions/{t.pk}/receipt.pdf")
        self.assertEqual(r["Content-Type"], "application/pdf")
        text = pdf_text(r.content)
        for expected in (t.receipt_no, t.reference, "40 000", "691 00 00 01", "MTN", "Cashier Fako 1"):
            self.assertIn(expected, text)
        self.assertNotIn("CORRECTED", text)
        m = self.client_for("manager")
        r = m.post(f"/transactions/{t.pk}/correct/", {"idem": str(uuid.uuid4()), "kind": "REPLACE",
                                                      "reason": "Was 45,000", "amount": "45000"})
        self.assertEqual(r.status_code, 302)
        text = pdf_text(c.get(f"/transactions/{t.pk}/receipt.pdf").content)
        self.assertIn("CORRECTED", text)
        self.assertIn("Was 45,000", text)
        self.assertContains(c.get(f"/transactions/{t.pk}/receipt/"), "CORRECTED")
        self.assertEqual(self.cbal(self.a1), D(-45000))
        fr = self.client_for("cashier.fako1", "fr")
        self.assertIn("CORRIGÉ", pdf_text(fr.get(f"/transactions/{t.pk}/receipt.pdf").content))

    def test_closing_flow_and_discrepancy_resolution_via_web(self):
        c = self.client_for("cashier.fako1")
        data = {f"acct_{a.id}": str(self.bal(a)) for a in closing_accounts(self.c_f1)}
        data[f"acct_{self.d_f1.id}"] = "990000"
        r = c.post("/closing/", data)  # without confirmation tick
        self.assertEqual(Closing.objects.count(), 0)
        data["confirm"] = "1"
        r = c.post("/closing/", data)
        self.assertEqual(r.status_code, 302)
        closing = Closing.objects.get()
        page = c.get(r["Location"])
        self.assertContains(page, "-10 000")
        d = Discrepancy.objects.get()
        self.assertEqual(c.post(f"/discrepancies/{d.pk}/", {"action": "resolve", "resolution": "COUNT_ERROR",
                                                            "notes": "x"}).status_code, 200)
        d.refresh_from_db()
        self.assertEqual(d.status, "OPEN")  # cashier cannot resolve
        m = self.client_for("finance")
        m.post(f"/discrepancies/{d.pk}/", {"action": "note", "text": "Checked CCTV"})
        m.post(f"/discrepancies/{d.pk}/", {"action": "resolve", "resolution": "ADJUSTMENT", "notes": "Write-off"})
        d.refresh_from_db()
        self.assertEqual(d.status, "RESOLVED")
        self.assertEqual(self.bal(self.d_f1), D(990000))
        self.assertTrue(closing.lines.filter(difference=D(-10000)).exists())
        rep = m.get("/reports/closings/?preset=all&format=pdf")
        self.assertIn("-10 000", pdf_text(rep.content))

    def test_reports_render_and_export_with_same_filters(self):
        ops.record_supply(self.c_f1, product="UVE", customer=self.a1, amount=D(123456), source_account=self.uve_f1)
        ops.record_supply(self.c_me, product="UVE", customer=self.a3, amount=D(654321), source_account=self.uve_me)
        m = self.client_for("manager")
        q = f"/reports/supplies/?preset=all&master_zone={self.zones['meme'].pk}"
        self.assertContains(m.get(q), "654 321")
        self.assertNotContains(m.get(q), "123 456")
        ws = load_workbook(io.BytesIO(m.get(q + "&format=xlsx").content)).active
        values = [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]
        self.assertIn(654321, values)
        self.assertNotIn(123456, values)
        text = pdf_text(m.get(q + "&format=pdf").content)
        self.assertIn("654 321", text)
        self.assertNotIn("123 456", text)
        rec = m.get("/reports/reconciliation/?preset=all")
        self.assertContains(rec, "Yes")
        self.assertNotContains(rec, "MISMATCH")


class ReconciliationColumns(WebBase):
    def test_overall_opening_is_zero_and_period_opening_carries(self):
        from datetime import timedelta

        from core.reports.defs import reconciliation
        from core.reports.engine import Filters
        from core.services.days import douala_today

        today = douala_today()
        rows = {r["account"]: r for r in reconciliation(Filters(None, None, user=self.manager)).tables[0].rows}
        r = rows[str(self.d_f1)]
        self.assertEqual((r["opening"], r["launch"], r["closing"]), (D(0), D(1000000), D(1000000)))
        nxt = today + timedelta(days=1)
        rows = {r["account"]: r for r in reconciliation(Filters(nxt, nxt, user=self.manager)).tables[0].rows}
        r = rows[str(self.d_f1)]
        self.assertEqual((r["opening"], r["launch"], r["closing"], r["check"]), (D(1000000), D(0), D(1000000), "OK"))
